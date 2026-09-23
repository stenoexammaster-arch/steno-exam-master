import os
import time
import uuid
import threading
import os
import pytesseract

pytesseract.pytesseract.tesseract_cmd = r"C:\Users\Asus\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"
os.environ["TESSDATA_PREFIX"] = r"C:\Users\Asus\AppData\Local\Programs\Tesseract-OCR\tessdata"
from pathlib import Path
from dataclasses import dataclass, field
from typing import Dict, Optional, List, Callable

from fastapi import FastAPI, Request, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from pydantic import BaseModel

import fitz  # pymupdf
import cv2
import numpy as np
from PIL import Image
import pytesseract
from docx import Document
from pypdf import PdfMerger


# ------------------- CONFIG -------------------
# If tesseract is not in PATH, set this (example):
# pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


# ------------------- PATHS -------------------
BASE_DIR = Path(__file__).resolve().parent.parent
UPLOADS_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "output"
LOGS_DIR = BASE_DIR / "logs"
TMP_DIR = UPLOADS_DIR / "tmp"

for d in (UPLOADS_DIR, OUTPUT_DIR, LOGS_DIR, TMP_DIR):
    d.mkdir(parents=True, exist_ok=True)


# ------------------- APP -------------------
app = FastAPI(title="Smart OCR Converter Backend", version="0.1")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],          # dev: allow any origin (file:// bhi)
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ------------------- MODELS -------------------
class UploadInitIn(BaseModel):
    filename: str
    size: int
    mime: Optional[str] = None


class UploadCompleteIn(BaseModel):
    uploadId: str


class ConvertIn(BaseModel):
    fileId: str
    options: dict


@dataclass
class UploadSession:
    upload_id: str
    filename: str
    size: int
    created_at: float = field(default_factory=time.time)
    tmp_dir: Path = None


@dataclass
class JobState:
    job_id: str
    file_id: str
    input_path: Path
    options: dict

    status: str = "queued"     # queued | running | completed | failed | cancelled | paused
    stage: str = "Queued"
    progress: float = 0.0
    eta_seconds: Optional[int] = None
    error: Optional[str] = None
    log: List[str] = field(default_factory=list)
    output_path: Optional[Path] = None

    pause_event: threading.Event = field(default_factory=threading.Event)
    cancel_event: threading.Event = field(default_factory=threading.Event)


UPLOADS: Dict[str, UploadSession] = {}
FILES: Dict[str, Path] = {}
JOBS: Dict[str, JobState] = {}


# ------------------- HELPERS -------------------
def safe_filename(name: str) -> str:
    name = name.replace("\\", "_").replace("/", "_").strip()
    return "".join(ch for ch in name if ch.isalnum() or ch in "._- ()")[:180] or "upload.bin"


def add_log(job: JobState, msg: str):
    ts = time.strftime("%H:%M:%S")
    job.log.append(f"{ts} {msg}")
    if len(job.log) > 200:
        job.log[:] = job.log[-200:]


def cleanup_file_later(path: Path, seconds: int = 1800):
    def _delete():
        try:
            if path.exists():
                path.unlink()
        except Exception:
            pass

    t = threading.Timer(seconds, _delete)
    t.daemon = True
    t.start()


def detect_pages(path: Path) -> Optional[int]:
    try:
        if path.suffix.lower() == ".pdf":
            doc = fitz.open(str(path))
            n = doc.page_count
            doc.close()
            return n
        return 1
    except Exception:
        return None


def tess_lang(language: str) -> str:
    if language == "eng":
        return "eng"
    if language == "hin":
        return "hin"
    if language == "mixed":
        return "eng+hin"
    return "eng"


def deskew(gray: np.ndarray) -> np.ndarray:
    inv = cv2.bitwise_not(gray)
    thresh = cv2.threshold(inv, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]
    coords = np.column_stack(np.where(thresh > 0))
    if coords.size == 0:
        return gray

    angle = cv2.minAreaRect(coords)[-1]
    if angle < -45:
        angle = -(90 + angle)
    else:
        angle = -angle

    h, w = gray.shape[:2]
    M = cv2.getRotationMatrix2D((w // 2, h // 2), angle, 1.0)
    return cv2.warpAffine(gray, M, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)


def preprocess(bgr: np.ndarray, preprocessing: dict) -> np.ndarray:
    denoise = bool(preprocessing.get("denoise", True))
    do_deskew = bool(preprocessing.get("deskew", True))
    contrast = bool(preprocessing.get("contrast", True))

    gray = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)

    if contrast:
        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        gray = clahe.apply(gray)

    if denoise:
        gray = cv2.fastNlMeansDenoising(gray, None, 20, 7, 21)

    if do_deskew:
        gray = deskew(gray)

    gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]
    return gray


def render_pdf_pages(pdf_path: Path, dpi: int = 300):
    doc = fitz.open(str(pdf_path))
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    for i in range(doc.page_count):
        page = doc.load_page(i)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        bgr = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
        yield (i + 1, doc.page_count, bgr)
    doc.close()


def load_image_pages(image_path: Path):
    pil = Image.open(str(image_path))
    frames = []
    try:
        i = 0
        while True:
            pil.seek(i)
            frames.append(pil.convert("RGB"))
            i += 1
    except EOFError:
        pass

    if frames:
        total = len(frames)
        for idx, frame in enumerate(frames, start=1):
            rgb = np.array(frame)
            bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
            yield (idx, total, bgr)
    else:
        rgb = np.array(pil.convert("RGB"))
        bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
        yield (1, 1, bgr)


def run_ocr(job: JobState, on_progress: Callable[[str, float, Optional[str]], None]) -> Path:
    opts = job.options or {}
    output = opts.get("output", "docx")
    language = opts.get("language", "eng")
    preprocessing = opts.get("preprocessing", {}) or {}
    target_dpi = opts.get("targetDpi", "300")

    dpi = 300
    if isinstance(target_dpi, str) and target_dpi.isdigit():
        dpi = int(target_dpi)

    lang = tess_lang(language)
    config = "--oem 1 --psm 6"

    is_pdf = job.input_path.suffix.lower() == ".pdf"
    pages = render_pdf_pages(job.input_path, dpi=dpi) if is_pdf else load_image_pages(job.input_path)

    base = f"{job.job_id}_{int(time.time())}"

    if output == "txt":
        out_path = OUTPUT_DIR / f"{base}.txt"
        chunks = []
        for pno, total, bgr in pages:
            if job.cancel_event.is_set():
                raise Exception("Cancelled")

            while job.pause_event.is_set():
                time.sleep(0.3)

            on_progress("Preprocessing", 5 + (pno / total) * 20, f"Page {pno}/{total}")
            gray = preprocess(bgr, preprocessing)

            on_progress("OCR", 25 + (pno / total) * 70, f"OCR page {pno}/{total}")
            text = pytesseract.image_to_string(gray, lang=lang, config=config)
            chunks.append(text.strip() + "\n")

        out_path.write_text("\n".join(chunks), encoding="utf-8")
        on_progress("Finalizing", 99, "TXT ready")
        return out_path

    if output == "pdf_searchable":
        out_path = OUTPUT_DIR / f"{base}.pdf"
        merger = PdfMerger()
        tmp_pages = []

        for pno, total, bgr in pages:
            if job.cancel_event.is_set():
                raise Exception("Cancelled")

            while job.pause_event.is_set():
                time.sleep(0.3)

            gray = preprocess(bgr, preprocessing)
            on_progress("OCR to Searchable PDF", 10 + (pno / total) * 85, f"Page {pno}/{total}")

            pdf_bytes = pytesseract.image_to_pdf_or_hocr(gray, lang=lang, extension="pdf")
            tmp_pdf = OUTPUT_DIR / f"{base}_p{pno:04d}.pdf"
            tmp_pdf.write_bytes(pdf_bytes)
            tmp_pages.append(tmp_pdf)
            merger.append(str(tmp_pdf))

        with out_path.open("wb") as f:
            merger.write(f)
        merger.close()

        for p in tmp_pages:
            try:
                p.unlink()
            except Exception:
                pass

        on_progress("Finalizing", 99, "Searchable PDF ready")
        return out_path

    # DOCX default
    out_path = OUTPUT_DIR / f"{base}.docx"
    doc = Document()

    for pno, total, bgr in pages:
        if job.cancel_event.is_set():
            raise Exception("Cancelled")

        while job.pause_event.is_set():
            time.sleep(0.3)

        on_progress("Preprocessing", 10 + (pno / total) * 10, f"Page {pno}/{total}")
        gray = preprocess(bgr, preprocessing)

        on_progress("OCR", 20 + (pno / total) * 75, f"OCR page {pno}/{total}")
        text = pytesseract.image_to_string(gray, lang=lang, config=config)

        # basic paragraph handling
        lines = [ln.rstrip() for ln in text.splitlines()]
        buff = []
        for ln in lines:
            if ln.strip() == "":
                if buff:
                    doc.add_paragraph(" ".join(buff).strip())
                    buff = []
            else:
                buff.append(ln.strip())
        if buff:
            doc.add_paragraph(" ".join(buff).strip())

        if pno != total:
            doc.add_page_break()

    on_progress("Finalizing", 99, "Writing DOCX")
    doc.save(str(out_path))
    on_progress("Completed", 100, "DOCX ready")
    return out_path


# ------------------- API -------------------
@app.post("/api/upload/init")
def upload_init(data: UploadInitIn):
    upload_id = str(uuid.uuid4())
    sess = UploadSession(
        upload_id=upload_id,
        filename=safe_filename(data.filename),
        size=int(data.size),
        tmp_dir=(TMP_DIR / upload_id),
    )
    sess.tmp_dir.mkdir(parents=True, exist_ok=True)
    UPLOADS[upload_id] = sess
    return {"uploadId": upload_id}


@app.post("/api/upload/chunk")
async def upload_chunk(
    request: Request,
    x_upload_id: str = Header(..., alias="X-Upload-Id"),
    x_chunk_index: int = Header(..., alias="X-Chunk-Index"),
    x_chunk_total: int = Header(..., alias="X-Chunk-Total"),
):
    sess = UPLOADS.get(x_upload_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Invalid uploadId")

    chunk_bytes = await request.body()
    if not chunk_bytes:
        raise HTTPException(status_code=400, detail="Empty chunk")

    chunk_path = sess.tmp_dir / f"chunk_{x_chunk_index:06d}.part"
    chunk_path.write_bytes(chunk_bytes)

    return JSONResponse({"ok": True, "chunkIndex": x_chunk_index, "chunkTotal": x_chunk_total})


@app.post("/api/upload/complete")
def upload_complete(data: UploadCompleteIn):
    sess = UPLOADS.get(data.uploadId)
    if not sess:
        raise HTTPException(status_code=404, detail="Invalid uploadId")

    file_id = str(uuid.uuid4())
    final_path = UPLOADS_DIR / f"{file_id}_{sess.filename}"

    chunk_files = sorted(sess.tmp_dir.glob("chunk_*.part"))
    if not chunk_files:
        raise HTTPException(status_code=400, detail="No chunks found")

    with final_path.open("wb") as out:
        for cf in chunk_files:
            out.write(cf.read_bytes())

    # cleanup tmp
    try:
        for cf in chunk_files:
            cf.unlink(missing_ok=True)
        sess.tmp_dir.rmdir()
    except Exception:
        pass

    FILES[file_id] = final_path
    UPLOADS.pop(data.uploadId, None)

    pages = detect_pages(final_path)
    cleanup_file_later(final_path, seconds=1800)

    return {"fileId": file_id, "pages": pages}


@app.post("/api/convert")
def convert(data: ConvertIn):
    input_path = FILES.get(data.fileId)
    if not input_path or not input_path.exists():
        raise HTTPException(status_code=404, detail="Invalid fileId")

    job_id = str(uuid.uuid4())
    job = JobState(job_id=job_id, file_id=data.fileId, input_path=input_path, options=data.options or {})
    JOBS[job_id] = job

    def worker():
        try:
            job.status = "running"
            job.stage = "Preparing"
            job.progress = 1.0
            add_log(job, "Job started")

            def on_progress(stage, prog, line=None):
                job.stage = stage
                job.progress = float(prog)
                if line:
                    add_log(job, line)

            out_path = run_ocr(job, on_progress)
            job.output_path = out_path
            job.status = "completed"
            job.stage = "Completed"
            job.progress = 100.0
            add_log(job, "Job completed")
            cleanup_file_later(out_path, seconds=1800)

        except Exception as e:
            job.status = "failed"
            job.stage = "Failed"
            job.error = str(e)
            add_log(job, f"ERROR: {e}")

    threading.Thread(target=worker, daemon=True).start()
    return {"jobId": job_id}


@app.get("/api/job/{job_id}/status")
def job_status(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Invalid jobId")

    download_url = None
    if job.status == "completed" and job.output_path:
        download_url = f"/api/download/{job_id}"

    log_tail = job.log[-1] if job.log else None

    return {
        "status": job.status,
        "stage": job.stage,
        "progress": job.progress,
        "etaSeconds": job.eta_seconds,
        "logTail": log_tail,
        "error": job.error,
        "downloadUrl": download_url,
    }


@app.get("/api/download/{job_id}")
def download(job_id: str):
    job = JOBS.get(job_id)
    if not job or job.status != "completed" or not job.output_path:
        raise HTTPException(status_code=404, detail="Output not ready")

    out_path = job.output_path
    if not out_path.exists():
        raise HTTPException(status_code=404, detail="Output file missing")

    return FileResponse(
        path=str(out_path),
        filename=out_path.name,
        media_type="application/octet-stream",
    )
@app.get("/api/health")
def health():
    return {"ok": True}

@app.get("/api/tesseract")
def tesseract_status():
    import os
    import pytesseract

    info = {
        "tesseract_cmd": getattr(pytesseract.pytesseract, "tesseract_cmd", None),
        "TESSDATA_PREFIX": os.environ.get("TESSDATA_PREFIX"),
        "version": None,
        "languages": None,
        "error": None,
    }

    try:
        info["version"] = str(pytesseract.get_tesseract_version())
        # list languages available
        try:
            info["languages"] = pytesseract.get_languages(config="")
        except Exception as e:
            info["languages"] = None
            info["error"] = f"Languages check failed: {e}"
    except Exception as e:
        info["error"] = f"Tesseract not working: {e}"

    return info